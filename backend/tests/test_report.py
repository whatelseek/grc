import io

from docx import Document

from app.models import Role


def _text_of(payload: bytes) -> str:
    document = Document(io.BytesIO(payload))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _assessment(client, headers, org, name="Аудит", items=()):
    body = client.post(
        "/api/assessments",
        json={"client_id": str(org.id), "target_ig": 1, "name": name},
        headers=headers,
    ).json()
    if items:
        client.patch(
            f"/api/assessments/{body['id']}/items", json={"items": list(items)}, headers=headers
        )
    return body


def test_report_covers_summary_plan_and_appendix(client, make_user, make_client_org, auth):
    org = make_client_org("Ромашка")
    consultant, _ = make_user(Role.consultant, clients=[org])
    headers = auth(consultant)
    assessment = _assessment(
        client,
        headers,
        org,
        "Аудит 2026",
        [
            {"safeguard_id": "1.1", "status": "implemented"},
            {
                "safeguard_id": "1.2",
                "status": "partial",
                "owner": "А. Петров",
                "due_date": "2026-12-01",
                "priority": "high",
            },
        ],
    )

    response = client.get(f"/api/assessments/{assessment['id']}/report.docx", headers=headers)
    assert response.status_code == 200
    assert "wordprocessingml" in response.headers["content-type"]
    assert ".docx" in response.headers["content-disposition"]

    text = _text_of(response.content)
    assert "Аудит мер защиты" in text
    assert "Ромашка" in text
    assert "План устранения" in text
    assert "Приложение. Все требования в объёме" in text
    # The partially implemented requirement carries its owner and date into the plan.
    assert "А. Петров" in text and "2026-12-01" in text
    assert "Высокий" in text
    # The implemented one is out of the plan but present in the appendix.
    assert "Реализация по контролям" in text
    assert "Реализация по функциям" in text


def test_report_states_how_many_ig1_gaps_remain(client, make_user, make_client_org, auth):
    org = make_client_org()
    consultant, _ = make_user(Role.consultant, clients=[org])
    headers = auth(consultant)
    assessment = _assessment(client, headers, org)

    text = _text_of(
        client.get(f"/api/assessments/{assessment['id']}/report.docx", headers=headers).content
    )
    # Nothing answered yet: all 56 IG1 requirements are open.
    assert "56 требований в объёме IG1" in text
    assert "Требуют работы 56 мер, из них 56 относятся к IG1" in text


def test_report_includes_the_dynamics_when_a_previous_audit_exists(
    client, make_user, make_client_org, auth
):
    org = make_client_org()
    consultant, _ = make_user(Role.consultant, clients=[org])
    headers = auth(consultant)
    _assessment(client, headers, org, "Прошлый аудит", [{"safeguard_id": "1.1", "status": "implemented"}])
    second = _assessment(
        client,
        headers,
        org,
        "Текущий аудит",
        [
            {"safeguard_id": "1.1", "status": "not_implemented"},
            {"safeguard_id": "1.2", "status": "implemented"},
        ],
    )

    text = _text_of(
        client.get(f"/api/assessments/{second['id']}/report.docx", headers=headers).content
    )
    assert "Динамика с прошлого аудита" in text
    assert "Прошлый аудит" in text and "Текущий аудит" in text
    # 1.1 went backwards and must be called out.
    assert "Откаты назад требуют объяснения" in text


def test_report_is_scoped_to_the_callers_clients(client, make_user, make_client_org, auth):
    mine, theirs = make_client_org("Мой"), make_client_org("Чужой")
    admin, _ = make_user(Role.admin)
    outsider, _ = make_user(Role.consultant, clients=[mine])
    assessment = _assessment(client, auth(admin), theirs)

    assert (
        client.get(f"/api/assessments/{assessment['id']}/report.docx", headers=auth(outsider)).status_code
        == 403
    )
